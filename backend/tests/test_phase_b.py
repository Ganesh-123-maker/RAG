from __future__ import annotations

import asyncio
import hashlib
import io
import uuid
from types import SimpleNamespace
from unittest.mock import AsyncMock, MagicMock, patch

import pytest
from fastapi import UploadFile
from sqlalchemy.ext.asyncio import AsyncSession

from app.services.agent_service import (
    _truncate_context_to_budget,
    _truncate_history_to_budget,
    reasoning_node,
)
from app.services.chunking_service import (
    ScannedDocumentError,
    chunk_text,
    extract_segments,
)
from app.services.document_service import DocumentService, sniff_mismatch
from app.services.rag_service import parse_citations
from app.services.vector_store_service import (
    CollectionDimensionMismatchError,
    IdentityReranker,
    RetrievedChunk,
    VectorStoreService,
    mmr_select,
    stable_chunk_id,
)


# ---------------------------------------------------------------- helpers
class _FakePage:
    def __init__(self, text: str, images: list | None = None) -> None:
        self._text = text
        self.images = images or []

    def extract_text(self) -> str:
        return self._text


class _FakePdf:
    def __init__(self, pages: list[_FakePage]) -> None:
        self.pages = pages

    def __enter__(self) -> _FakePdf:  # noqa: PYI034 -- test fake, not a stub
        return self

    def __exit__(self, *args: object) -> bool:
        return False


def _make_docx_bytes() -> bytes:
    from docx import Document as DocxDocument

    doc = DocxDocument()
    doc.add_heading("Findings", level=1)
    doc.add_paragraph("Key result under findings.")
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Metric"
    table.rows[0].cells[1].text = "Value"
    table.rows[1].cells[0].text = "Accuracy"
    table.rows[1].cells[1].text = "0.91"
    doc.add_heading("Risks", level=1)
    doc.add_paragraph("Risk paragraph.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- B1
class TestExtractionMetadata:
    def test_pdf_page_numbers(self) -> None:
        pages = [_FakePage("Page one text."), _FakePage("Page two text.")]
        with patch("pdfplumber.open", return_value=_FakePdf(pages)):
            segments = extract_segments(b"%PDF-fake", "pdf")
        assert [s.page for s in segments] == [1, 2]
        assert segments[0].text == "Page one text."

    def test_scanned_pdf_raises(self) -> None:
        pages = [_FakePage("", images=[object()]), _FakePage("", images=[object()])]
        with patch("pdfplumber.open", return_value=_FakePdf(pages)), pytest.raises(
            ScannedDocumentError, match="scanned"
        ):
            extract_segments(b"%PDF-fake", "pdf")

    def test_docx_tables_and_headings(self) -> None:
        segments = extract_segments(_make_docx_bytes(), "docx")
        texts = [s.text for s in segments]
        assert any("Metric | Value" in t for t in texts), texts
        assert any("Accuracy | 0.91" in t for t in texts), texts
        findings = next(s for s in segments if "Key result" in s.text)
        assert findings.section == "Findings"
        risks = next(s for s in segments if "Risk paragraph" in s.text)
        assert risks.section == "Risks"

    def test_md_headings(self) -> None:
        md = b"# Intro\n\nHello world.\n\n## Details\n\nMore detail here.\n"
        segments = extract_segments(md, "md")
        detail = next(s for s in segments if "More detail" in s.text)
        assert detail.section == "Details"

    def test_sniff_mismatch(self) -> None:
        assert sniff_mismatch(b"plain text here", "pdf") is True
        assert sniff_mismatch(b"%PDF-1.4", "pdf") is False
        assert sniff_mismatch(b"plain text", "docx") is True
        assert sniff_mismatch(b"PK\x03\x04", "docx") is False
        assert sniff_mismatch(b"%PDF-1.4", "txt") is True
        assert sniff_mismatch(b"plain text", "txt") is False


@pytest.mark.asyncio
async def test_magic_mismatch_returns_422(
    db_session: AsyncSession, fake_user, tmp_path, monkeypatch
) -> None:
    import app.services.document_service as doc_svc_mod

    monkeypatch.setattr(doc_svc_mod.settings, "UPLOAD_DIR", str(tmp_path))
    svc = DocumentService(db_session)
    # PDF magic missing despite a .pdf name.
    bad = UploadFile(filename="fake.pdf", file=io.BytesIO(b"not a pdf at all"))
    with pytest.raises(Exception) as exc:
        await svc.ingest(bad, fake_user.id)
    assert exc.value.status_code == 422
    assert list(tmp_path.rglob("*")) == []


@pytest.mark.asyncio
async def test_scanned_pdf_service_error(
    db_session: AsyncSession, fake_user, tmp_path, monkeypatch
) -> None:
    import app.services.document_service as doc_svc_mod

    monkeypatch.setattr(doc_svc_mod.settings, "UPLOAD_DIR", str(tmp_path))
    svc = DocumentService(db_session)
    pages = [_FakePage("", images=[object()])]
    with (
        patch("pdfplumber.open", return_value=_FakePdf(pages)),
        pytest.raises(Exception) as exc,
    ):
        bad = UploadFile(filename="scan.pdf", file=io.BytesIO(b"%PDF-1.4 scanned"))
        await svc.ingest(bad, fake_user.id)
    assert exc.value.status_code == 422
    assert "scanned" in str(exc.value.detail).lower()
    assert list(tmp_path.rglob("*")) == []


class _TrackingBytesIO(io.BytesIO):
    """BytesIO that records every read size and forbids unbounded reads."""

    def __init__(self, *args: object, **kwargs: object) -> None:
        super().__init__(*args, **kwargs)
        self.max_read = 0

    def read(self, size: int = -1) -> bytes:  # type: ignore[override]
        assert size is not None and size > 0, "unbounded read attempted"
        assert size <= 1024 * 1024, f"read size {size} exceeds 1 MiB bound"
        self.max_read = max(self.max_read, size)
        return super().read(size)


@pytest.mark.asyncio
async def test_oversize_rejected_without_full_load(
    db_session: AsyncSession, fake_user, tmp_path, monkeypatch
) -> None:
    import app.services.document_service as doc_svc_mod

    monkeypatch.setattr(doc_svc_mod.settings, "UPLOAD_DIR", str(tmp_path))
    monkeypatch.setattr(doc_svc_mod.settings, "MAX_UPLOAD_SIZE_MB", 50)
    svc = DocumentService(db_session)
    tracking = _TrackingBytesIO(b"x" * (51 * 1024 * 1024))
    big = UploadFile(filename="big.txt", file=tracking)
    with pytest.raises(Exception) as exc:
        await svc.ingest(big, fake_user.id)
    assert exc.value.status_code == 413
    assert tracking.max_read <= 1024 * 1024
    assert list(tmp_path.rglob("*")) == []


@pytest.mark.asyncio
async def test_parsing_runs_in_threadpool(
    db_session: AsyncSession, fake_user, tmp_path, monkeypatch
) -> None:
    import app.services.document_service as doc_svc_mod

    monkeypatch.setattr(doc_svc_mod.settings, "UPLOAD_DIR", str(tmp_path))
    svc = DocumentService(db_session)
    real_to_thread = asyncio.to_thread
    called: list[str] = []

    async def counting_to_thread(func, *args, **kwargs):  # type: ignore[no-untyped-def]
        called.append(getattr(func, "__name__", str(func)))
        return await real_to_thread(func, *args, **kwargs)

    with (
        patch.object(
            svc._embed_svc,
            "embed_batch",
            new_callable=AsyncMock,
            side_effect=lambda texts: [[0.1] * 3072 for _ in texts],
        ),
        patch.object(svc._vector_svc, "ensure_collection", new_callable=AsyncMock),
        patch.object(svc._vector_svc, "upsert_chunks", new_callable=AsyncMock),
        patch("asyncio.to_thread", side_effect=counting_to_thread),
    ):
        file = UploadFile(
            filename="doc.txt", file=io.BytesIO(b"Threadpool test content here.")
        )
        await svc.ingest(file, fake_user.id)
    assert "extract_segments" in called
    assert "chunk_text" in called


# ---------------------------------------------------------------- B2
class TestChunking:
    def test_sentence_boundaries_offsets_and_hash(self) -> None:
        text = (
            "Alpha one two three four. Beta five six seven eight nine. "
            "Gamma ten eleven twelve thirteen fourteen fifteen."
        )
        doc_id = uuid.uuid4()
        chunks = chunk_text(text, doc_id, "doc.txt", chunk_size=16, overlap=4)
        assert len(chunks) >= 2
        for chunk in chunks:
            assert text[chunk.char_start : chunk.char_end] == chunk.text
            assert chunk.text.strip().endswith(".")
            assert (
                chunk.content_hash
                == hashlib.sha256(chunk.text.encode("utf-8")).hexdigest()
            )
        starts = [c.char_start for c in chunks]
        assert starts == sorted(starts)

    def test_chunk_page_and_section_metadata(self) -> None:
        from app.services.chunking_service import ExtractedSegment

        segments = [
            ExtractedSegment(text="First page sentence.", page=1, section="Intro"),
            ExtractedSegment(text="Second page sentence.", page=2, section="Body"),
        ]
        raw = "\n\n".join(s.text for s in segments)
        chunks = chunk_text(
            raw,
            uuid.uuid4(),
            "doc.pdf",
            chunk_size=512,
            overlap=64,
            segments=segments,
        )
        assert chunks[0].page == 1
        assert chunks[0].section == "Intro"


# ---------------------------------------------------------------- B3
@pytest.mark.asyncio
async def test_document_dedupe_same_owner_and_other_owner(
    db_session: AsyncSession, fake_user, tmp_path, monkeypatch
) -> None:
    import app.services.document_service as doc_svc_mod
    from app.models.user import User
    from app.repositories.document_repository import DocumentRepository

    monkeypatch.setattr(doc_svc_mod.settings, "UPLOAD_DIR", str(tmp_path))
    other = User(
        id=uuid.uuid4(),
        email="other@example.com",
        hashed_password="pw",
        full_name="Other",
    )
    db_session.add(other)
    await db_session.flush()

    svc = DocumentService(db_session)
    content = b"Duplicate document content for dedupe testing."
    with (
        patch.object(
            svc._embed_svc,
            "embed_batch",
            new_callable=AsyncMock,
            side_effect=lambda texts: [[0.1] * 3072 for _ in texts],
        ),
        patch.object(svc._vector_svc, "ensure_collection", new_callable=AsyncMock),
        patch.object(
            svc._vector_svc, "upsert_chunks", new_callable=AsyncMock
        ) as mock_upsert,
    ):
        first = await svc.ingest(
            UploadFile(filename="dup.txt", file=io.BytesIO(content)), fake_user.id
        )
        second = await svc.ingest(
            UploadFile(filename="dup.txt", file=io.BytesIO(content)), fake_user.id
        )
        assert second == first
        assert mock_upsert.await_count == 1
        third = await svc.ingest(
            UploadFile(filename="dup.txt", file=io.BytesIO(content)), other.id
        )
        assert third != first
        assert mock_upsert.await_count == 2

    repo = DocumentRepository(db_session)
    assert len(await repo.get_by_owner(fake_user.id)) == 1
    assert len(await repo.get_by_owner(other.id)) == 1


@pytest.mark.asyncio
async def test_chunk_dedupe_within_document(
    db_session: AsyncSession, fake_user, tmp_path, monkeypatch
) -> None:
    import app.services.document_service as doc_svc_mod

    monkeypatch.setattr(doc_svc_mod.settings, "UPLOAD_DIR", str(tmp_path))
    monkeypatch.setattr(doc_svc_mod.settings, "CHUNK_SIZE_TOKENS", 4)
    monkeypatch.setattr(doc_svc_mod.settings, "CHUNK_OVERLAP_TOKENS", 0)
    svc = DocumentService(db_session)
    captured: list[list] = []
    content = b"One two three. One two three. Four five six."
    with (
        patch.object(
            svc._embed_svc,
            "embed_batch",
            new_callable=AsyncMock,
            side_effect=lambda texts: [[0.1] * 3072 for _ in texts],
        ),
        patch.object(svc._vector_svc, "ensure_collection", new_callable=AsyncMock),
        patch.object(
            svc._vector_svc,
            "upsert_chunks",
            new_callable=AsyncMock,
            side_effect=lambda chunks, embeddings, owner: captured.append(chunks),
        ),
    ):
        doc_id = await svc.ingest(
            UploadFile(filename="rep.txt", file=io.BytesIO(content)), fake_user.id
        )
    assert len(captured) == 1
    assert len(captured[0]) == 2  # identical sentence chunk kept once
    assert [c.chunk_index for c in captured[0]] == [0, 1]
    repo = doc_svc_mod.DocumentRepository(db_session)
    doc = await repo.get_by_id_for_owner(doc_id, fake_user.id)
    assert doc is not None and doc.chunk_count == 2


# ---------------------------------------------------------------- B4
class TestRetrieval:
    def test_mmr_diversity(self) -> None:
        def chunk(score: float, vector: list[float]) -> RetrievedChunk:
            return RetrievedChunk(
                text="t",
                document_id="d",
                document_name="d.txt",
                chunk_index=0,
                score=score,
                vector=vector,
            )

        candidates = [
            chunk(0.95, [1.0, 0.0]),
            chunk(0.94, [0.99, 0.01]),
            chunk(0.80, [0.0, 1.0]),
        ]
        selected = mmr_select(candidates, top_k=2, lambda_mult=0.7)
        assert [c.score for c in selected] == [0.95, 0.80]

    def test_identity_reranker_keeps_order(self) -> None:
        chunks = [
            RetrievedChunk(
                text="a",
                document_id="d",
                document_name="d.txt",
                chunk_index=0,
                score=0.9,
            ),
            RetrievedChunk(
                text="b",
                document_id="d",
                document_name="d.txt",
                chunk_index=1,
                score=0.8,
            ),
        ]
        assert IdentityReranker().rerank("q", chunks) == chunks

    @pytest.mark.asyncio
    async def test_threshold_filters_low_scores(self) -> None:
        svc = VectorStoreService()
        owner_id = uuid.uuid4()
        points = []
        for score in (0.9, 0.2, 0.05):
            point = MagicMock()
            point.score = score
            point.vector = None
            point.payload = {
                "text": f"chunk {score}",
                "document_id": "d",
                "document_name": "d.txt",
                "chunk_index": 0,
            }
            points.append(point)
        with patch.object(
            svc._client, "search", new_callable=AsyncMock, return_value=points
        ):
            results = await svc.search(
                query_vector=[0.1] * 3072, owner_id=owner_id, threshold=0.30
            )
        assert [r.score for r in results] == [0.9]

    @pytest.mark.asyncio
    async def test_collection_dimension_mismatch_raises(self) -> None:
        svc = VectorStoreService()
        collections = MagicMock()
        collections.collections = [MagicMock(name="rag_chunks")]
        # The name mock above does not carry a real .name string in all
        # mock configurations; set it explicitly for this test.
        collections.collections[0].name = "rag_chunks"
        info = SimpleNamespace(
            config=SimpleNamespace(
                params=SimpleNamespace(vectors=SimpleNamespace(size=1024))
            )
        )
        with (
            patch.object(
                svc._client,
                "get_collections",
                new_callable=AsyncMock,
                return_value=collections,
            ),
            patch.object(
                svc._client, "get_collection", new_callable=AsyncMock, return_value=info
            ),
            pytest.raises(
                CollectionDimensionMismatchError, match="1024.*3072|3072.*1024"
            ),
        ):
            await svc.ensure_collection()

    @pytest.mark.asyncio
    async def test_upsert_payload_carries_metadata(self) -> None:
        from app.services.chunking_service import TextChunk

        svc = VectorStoreService()
        owner_id = uuid.uuid4()
        doc_id = uuid.uuid4()
        chunk = TextChunk(
            text="Hello world.",
            chunk_index=0,
            document_id=doc_id,
            document_name="doc.pdf",
            page=3,
            section="Intro",
        )
        with patch.object(svc._client, "upsert", new_callable=AsyncMock) as mock_upsert:
            await svc.upsert_chunks([chunk], [[0.1] * 3072], owner_id)
        point = mock_upsert.call_args.kwargs["points"][0]
        assert point.payload["page"] == 3
        assert point.payload["section"] == "Intro"
        assert point.payload["chunk_id"] == stable_chunk_id(owner_id, chunk)
        assert point.payload["content_hash"] == chunk.content_hash
        # Same chunk upserted twice collapses to one point.
        with patch.object(
            svc._client, "upsert", new_callable=AsyncMock
        ) as mock_upsert2:
            await svc.upsert_chunks([chunk, chunk], [[0.1] * 3072] * 2, owner_id)
        assert len(mock_upsert2.call_args.kwargs["points"]) == 1


@pytest.mark.asyncio
async def test_empty_retrieval_refuses() -> None:
    state = {
        "question": "What is in the docs?",
        "owner_id": str(uuid.uuid4()),
        "document_id": None,
        "history": [],
        "retrieval_query": "What is in the docs?",
        "chunks": [],
        "answer": "",
        "citations": [],
        "needs_retrieval": True,
    }
    result = await reasoning_node(state)
    assert "couldn't find" in result["answer"]
    assert result["citations"] == []


# ---------------------------------------------------------------- B5
class TestCitationsAndBudgets:
    def test_citation_id_validation(self) -> None:
        doc_id = str(uuid.uuid4())
        chunks = [
            RetrievedChunk(
                text="Real chunk text.",
                document_id=doc_id,
                document_name="real.txt",
                chunk_index=0,
                score=0.9,
                chunk_id="chunk-abc",
                page=2,
                section="Intro",
            )
        ]
        _raw, cits = parse_citations(
            "Answer.\n\nCITATIONS:\n[chunk-abc] real.txt, chunk 0\n"
            "[unknown-id] other.txt, chunk 5\n[chunk-abc] real.txt, chunk 0",
            chunks,
        )
        assert len(cits) == 1
        assert cits[0]["chunk_id"] == "chunk-abc"
        assert cits[0]["page"] == 2
        assert cits[0]["section"] == "Intro"

    def test_positional_citation_is_not_substituted(self) -> None:
        chunks = [
            RetrievedChunk(
                text="Only chunk.",
                document_id="d",
                document_name="d.txt",
                chunk_index=0,
                score=0.9,
                chunk_id="real-id-1",
            )
        ]
        _, cits = parse_citations("Answer.\n\nCITATIONS:\n[1] d.txt, chunk 0", chunks)
        assert cits == []

    def test_context_budget_drops_lowest_score_first(self) -> None:
        chunks = [
            {"text": "word " * 100, "score": 0.99},
            {"text": "word " * 100, "score": 0.50},
            {"text": "word " * 100, "score": 0.10},
        ]
        kept = _truncate_context_to_budget(chunks, max_tokens=150)
        scores = [c["score"] for c in kept]
        assert 0.10 not in scores
        assert 0.99 in scores

    def test_history_budget_keeps_most_recent(self) -> None:
        history = [
            {"role": "user", "content": "old " * 100},
            {"role": "assistant", "content": "mid " * 10},
            {"role": "user", "content": "new question"},
        ]
        kept = _truncate_history_to_budget(history, max_tokens=50)
        assert kept[-1]["content"] == "new question"
        assert all("old " not in str(m["content"]) for m in kept)
